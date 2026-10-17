"""
The hand check: the engine recomputed in a spreadsheet
======================================================
    python -m src.engine.handcheck      # rebuilds reports/phase5_hand_check.xlsx

The workbook holds two years of simulated daily returns for two instruments and recomputes every
step of the engine with plain Excel formulas, one column per step: the mix under each recipe,
monthly RV, the signal, c_t over expanding ranges, the capped size, the positions, drift, turnover,
trading cost, borrowing, the short fee, net and excess returns, for the timed book and for
buy-and-hold. Sheet "Check" lines up five holding months against the engine.

The five months cover the places indexing goes wrong:
    2016-08  the first month with a size: the entry trade from cash; capped; borrows cash
    2016-10  below the cap, no borrowing
    2016-12  the last month of recipe 1; capped
    2017-01  the recipe changes: the size uses December under the new recipe, and A turns short
    2017-12  the last month
Planting a bug confirms the check bites: c_t using a pair one month early, sizing without the
one-month lag, ignoring drift, trading a day late, sizing January on the old recipe, dropping
financing, or needing 7 pairs instead of 6 each moves some number by 4e-4 or more (a clean run
differs by about 2e-14).

The case is set up to exercise every part of the engine: recipe 1 borrows cash (positions sum to
1.4), recipe 2 shorts instrument A, the cap is 1.3, and c_t needs only 6 pairs so that two years
are enough.

Building needs xlsxwriter and the `formulas` package (an Excel formula engine), which computes
every formula and stores its value in the file. test_engine_checks.py reads those stored values
with openpyxl and compares them with a fresh engine run, so the test needs neither. Excel or
Numbers recompute the formulas on opening.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

from src import engine as e
from src.config import REPORTS_DIR
from src.engine import synthetic

PATH = REPORTS_DIR / "phase5_hand_check.xlsx"
PARAMS = {"max_size": 1.3, "min_months": 6, "kappa_A": 5e-4, "kappa_B": 10e-4, "borrow_spread": 0.005,
          "short_fee_A": 0.0025, "short_fee_B": 0.0}
RECIPES = pd.DataFrame({"A": [0.8, -0.3], "B": [0.6, 1.2]}, index=pd.Index([2016, 2017], name="year"))
CHECK_MONTHS = ["2016-08", "2016-10", "2016-12", "2017-01", "2017-12"]
# Exchange holidays in 2016-2017 (NYSE), so the calendar has real gaps.
HOLIDAYS = ["2016-01-01", "2016-01-18", "2016-02-15", "2016-03-25", "2016-05-30", "2016-07-04", "2016-09-05",
            "2016-11-24", "2016-12-26", "2017-01-02", "2017-01-16", "2017-02-20", "2017-04-14", "2017-05-29",
            "2017-07-04", "2017-09-04", "2017-11-23", "2017-12-25"]
QUANTITIES = ["forecast", "c", "size", "w_A", "w_B", "drift_A", "drift_B", "turnover", "trading", "net", "excess",
              "bench_turnover", "bench_net", "bench_excess"]


def inputs() -> pd.DataFrame:
    """Daily excess returns of A and B and a rising risk-free rate, 2016-2017."""
    days = pd.bdate_range("2016-01-01", "2017-12-31", name="date").difference(pd.DatetimeIndex(HOLIDAYS))
    a = synthetic.garch(years=2.2, vol=0.18, seed=11)["x"].to_numpy()[:len(days)]
    b = synthetic.garch(years=2.2, vol=0.12, seed=12)["x"].to_numpy()[:len(days)]
    rf = np.linspace(0.0025, 0.0125, len(days)) / 252
    return pd.DataFrame({"A": a, "B": b, "rf": rf}, index=days)


def engine_run(daily: pd.DataFrame, params: dict = PARAMS, recipes: pd.DataFrame = RECIPES) -> pd.DataFrame:
    """The engine on the hand-check case: the baseline rule (inverse variance on last month's RV)."""
    months = daily.index.to_period("M").unique()
    positions = recipes.loc[months.year].set_axis(months)
    book = e.Book(daily[["A", "B"]], daily["rf"], positions)
    costs = e.Costs(kappa={"A": params["kappa_A"], "B": params["kappa_B"]}, borrow_spread=params["borrow_spread"],
                    short_fee={"A": params["short_fee_A"], "B": params["short_fee_B"]})
    return e.run(book, "inv_var", max_size=params["max_size"], costs=costs, min_months=params["min_months"])


# ---------------------------------------------------------------- the workbook

def _col(n: int) -> str:
    s = ""
    while n >= 0:
        s, n = chr(65 + n % 26) + s, n // 26 - 1
    return s


def _cells(daily: pd.DataFrame) -> dict:
    """Every cell of the workbook: {(sheet, "A1"): value or "=formula"}."""
    cells = {}
    put = lambda sheet, ref, x: cells.__setitem__((sheet, ref), x)

    # Inputs: parameters and recipes
    put("Inputs", "A1", "parameter"), put("Inputs", "B1", "value")
    names = list(PARAMS) + ["days_per_year"]
    for i, k in enumerate(names, start=2):
        put("Inputs", f"A{i}", k), put("Inputs", f"B{i}", PARAMS.get(k, 252))
    P = {k: f"Inputs!$B${i}" for i, k in enumerate(names, start=2)}
    put("Inputs", "D1", "recipe year (held in)"), put("Inputs", "E1", "position A"), put("Inputs", "F1", "position B")
    for i, (year, row) in enumerate(RECIPES.iterrows(), start=2):
        put("Inputs", f"D{i}", year), put("Inputs", f"E{i}", row["A"]), put("Inputs", f"F{i}", row["B"])
    last_recipe = 1 + len(RECIPES)
    pos_a, pos_b, years = f"Inputs!$E$2:$E${last_recipe}", f"Inputs!$F$2:$F${last_recipe}", f"Inputs!$D$2:$D${last_recipe}"

    # Daily: inputs, the mix under each recipe, and log growth terms
    heads = ["date", "month", "x_A", "x_B", "rf", "p1", "p2", "p1^2", "p2^2", "ln(1+p1+rf)", "ln(1+p2+rf)",
             "ln(1+rf)", "ln(1+x_A+rf)", "ln(1+x_B+rf)", "ln(1+x_A+rf+fee_A)", "ln(1+x_B+rf+fee_B)", "ln(1+rf+spread)"]
    for j, h in enumerate(heads):
        put("Daily", f"{_col(j)}1", h)
    for i, (d, row) in enumerate(daily.iterrows(), start=2):
        put("Daily", f"A{i}", d.to_pydatetime()), put("Daily", f"B{i}", d.strftime("%Y-%m"))
        put("Daily", f"C{i}", row["A"]), put("Daily", f"D{i}", row["B"]), put("Daily", f"E{i}", row["rf"])
        put("Daily", f"F{i}", f"=Inputs!$E$2*C{i}+Inputs!$F$2*D{i}")
        put("Daily", f"G{i}", f"=Inputs!$E$3*C{i}+Inputs!$F$3*D{i}")
        put("Daily", f"H{i}", f"=F{i}^2"), put("Daily", f"I{i}", f"=G{i}^2")
        put("Daily", f"J{i}", f"=LN(1+F{i}+E{i})"), put("Daily", f"K{i}", f"=LN(1+G{i}+E{i})")
        put("Daily", f"L{i}", f"=LN(1+E{i})")
        put("Daily", f"M{i}", f"=LN(1+C{i}+E{i})"), put("Daily", f"N{i}", f"=LN(1+D{i}+E{i})")
        put("Daily", f"O{i}", f"=LN(1+C{i}+E{i}+{P['short_fee_A']}/{P['days_per_year']})")
        put("Daily", f"P{i}", f"=LN(1+D{i}+E{i}+{P['short_fee_B']}/{P['days_per_year']})")
        put("Daily", f"Q{i}", f"=LN(1+E{i}+{P['borrow_spread']}/{P['days_per_year']})")
    n = len(daily) + 1
    D = lambda c: f"Daily!${c}$2:${c}${n}"
    sumifs = lambda c, r: f"SUMIFS({D(c)},{D('B')},A{r})"

    # Monthly: one row per calendar month
    months = daily.index.to_period("M").unique()
    heads = {"A": "month", "B": "recipe held", "C": "year", "D": "Rf (month)",
             "E": "ret1: excess return, mix 1", "F": "RV1", "G": "signal1 = 1/RV1", "H": "y1 = next ret1",
             "I": "z1 = y1 x signal1", "J": "c1 (pairs up to the row before)",
             "K": "ret2", "L": "RV2", "M": "signal2", "N": "y2", "O": "z2", "P": "c2",
             "Q": "recipe next month", "R": "forecast (RV of the next recipe's mix)", "S": "signal",
             "T": "c", "U": "c x signal", "V": "size for next month (capped)",
             "W": "size held", "X": "w_A", "Y": "w_B", "Z": "cash", "AA": "growth A", "AB": "growth B",
             "AC": "growth cash", "AD": "account growth", "AE": "drift_A", "AF": "drift_B", "AG": "|trade A|",
             "AH": "|trade B|", "AI": "turnover", "AJ": "trading cost", "AK": "net", "AL": "excess",
             "AM": "bench w_A", "AN": "bench w_B", "AO": "bench cash", "AP": "bench growth A", "AQ": "bench growth B",
             "AR": "bench growth cash", "AS": "bench account growth", "AT": "bench drift_A", "AU": "bench drift_B",
             "AV": "bench |trade A|", "AW": "bench |trade B|", "AX": "bench turnover", "AY": "bench trading cost",
             "AZ": "bench net", "BA": "bench excess"}
    for c, h in heads.items():
        put("Monthly", f"{c}1", h)
    last = len(months) + 1
    have_c = lambda r: r - 2 >= PARAMS["min_months"]           # rows with enough pairs above them
    held_rows = [r for r in range(3, last + 1) if have_c(r - 1)]
    for r, m in enumerate(months, start=2):
        put("Monthly", f"A{r}", str(m))
        put("Monthly", f"C{r}", f"=VALUE(LEFT(A{r},4))")
        put("Monthly", f"B{r}", f"=MATCH(C{r},{years},0)")
        put("Monthly", f"D{r}", f"=EXP({sumifs('L', r)})-1")
        for ret, rv, sig, y, z, c, lncol, sqcol in (("E", "F", "G", "H", "I", "J", "J", "H"),
                                                    ("K", "L", "M", "N", "O", "P", "K", "I")):
            put("Monthly", f"{ret}{r}", f"=EXP({sumifs(lncol, r)})-EXP({sumifs('L', r)})")
            put("Monthly", f"{rv}{r}", f"={sumifs(sqcol, r)}")
            put("Monthly", f"{sig}{r}", f"=1/{rv}{r}")
            if r < last:
                put("Monthly", f"{y}{r}", f"={ret}{r + 1}")
                put("Monthly", f"{z}{r}", f"={y}{r}*{sig}{r}")
            if have_c(r):
                rng = lambda col: f"{col}$2:{col}{r - 1}"
                put("Monthly", f"{c}{r}", f"=STDEV({rng(y)})/STDEV({rng(z)})")
        if r < last and have_c(r):                               # a decision for next month
            put("Monthly", f"Q{r}", f"=B{r + 1}")
            put("Monthly", f"R{r}", f"=IF(Q{r}=1,F{r},L{r})")
            put("Monthly", f"S{r}", f"=IF(Q{r}=1,G{r},M{r})")
            put("Monthly", f"T{r}", f"=IF(Q{r}=1,J{r},P{r})")
            put("Monthly", f"U{r}", f"=T{r}*S{r}")
            put("Monthly", f"V{r}", f"=MIN({P['max_size']},MAX(0,U{r}))")
        if r in held_rows:                                       # a holding month
            first = r == held_rows[0]
            for pre, size, wa, wb, cash, ga, gb, gc, acct, da, db, ta, tb, turn, cost, net, exc in (
                    ("", f"V{r - 1}", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL"),
                    ("bench", "1", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AU", "AV", "AW", "AX", "AY", "AZ", "BA")):
                if not pre:
                    put("Monthly", f"W{r}", f"={size}")
                    size = f"W{r}"
                put("Monthly", f"{wa}{r}", f"={size}*INDEX({pos_a},B{r})")
                put("Monthly", f"{wb}{r}", f"={size}*INDEX({pos_b},B{r})")
                put("Monthly", f"{cash}{r}", f"=1-{wa}{r}-{wb}{r}")
                put("Monthly", f"{ga}{r}", f"=IF({wa}{r}<0,EXP({sumifs('O', r)}),EXP({sumifs('M', r)}))")
                put("Monthly", f"{gb}{r}", f"=IF({wb}{r}<0,EXP({sumifs('P', r)}),EXP({sumifs('N', r)}))")
                put("Monthly", f"{gc}{r}", f"=IF({cash}{r}<0,EXP({sumifs('Q', r)}),1+D{r})")
                put("Monthly", f"{acct}{r}", f"={wa}{r}*{ga}{r}+{wb}{r}*{gb}{r}+{cash}{r}*{gc}{r}")
                put("Monthly", f"{da}{r}", f"={wa}{r}*{ga}{r}/{acct}{r}")
                put("Monthly", f"{db}{r}", f"={wb}{r}*{gb}{r}/{acct}{r}")
                put("Monthly", f"{ta}{r}", f"=ABS({wa}{r}-{0 if first else f'{da}{r - 1}'})")
                put("Monthly", f"{tb}{r}", f"=ABS({wb}{r}-{0 if first else f'{db}{r - 1}'})")
                put("Monthly", f"{turn}{r}", f"={ta}{r}+{tb}{r}")
                put("Monthly", f"{cost}{r}", f"={P['kappa_A']}*{ta}{r}+{P['kappa_B']}*{tb}{r}")
                put("Monthly", f"{net}{r}", f"=(1-{cost}{r})*{acct}{r}-1")
                put("Monthly", f"{exc}{r}", f"={net}{r}-D{r}")

    # Check: the spreadsheet against the engine, five months
    row_of = {str(m): r for r, m in enumerate(months, start=2)}
    source = {"forecast": ("R", -1), "c": ("T", -1), "size": ("V", -1), "w_A": ("X", 0), "w_B": ("Y", 0),
              "drift_A": ("AE", 0), "drift_B": ("AF", 0), "turnover": ("AI", 0), "trading": ("AJ", 0),
              "net": ("AK", 0), "excess": ("AL", 0), "bench_turnover": ("AX", 0), "bench_net": ("AZ", 0),
              "bench_excess": ("BA", 0)}
    for j, h in enumerate(["holding month", "quantity", "spreadsheet", "engine", "difference"]):
        put("Check", f"{_col(j)}1", h)
    r = 2
    for m in CHECK_MONTHS:
        for q in QUANTITIES:
            col, lag = source[q]
            put("Check", f"A{r}", m), put("Check", f"B{r}", q)
            put("Check", f"C{r}", f"=Monthly!{col}{row_of[m] + lag}")
            put("Check", f"E{r}", f"=C{r}-D{r}")
            r += 1
    return cells


NOTES = [
    "Phase 5 hand check: the backtest engine recomputed with plain spreadsheet formulas.",
    "Built by `python -m src.engine.handcheck`; test/test_engine_checks.py compares the stored formula values with a fresh engine run.",
    "",
    "Inputs: simulated daily excess returns of two instruments, A and B, and a daily risk-free rate (sheet Daily, columns A-E).",
    "Two recipes: positions per unit of size are held in every month of their year (sheet Inputs).",
    "Rule: size for month m+1 = min(cap, c_t / RV_t), with RV_t and c_t from the mix held in m+1, at the end of month t = m.",
    "c_t = STDEV(next-month returns) / STDEV(next-month return x signal), over the pairs complete by the end of t (expanding range).",
    "Execution: trade at the month's last close; positions drift within the month; turnover = |new - drifted|, cost = kappa x turnover.",
    "Cash below zero grows at rf + borrow spread; a short grows at its return + short fee (both accrued daily, 252 a year).",
    "Net return = (1 - trading cost) x account growth - 1. Excess = net - Rf of the month.",
    "Buy-and-hold: the same recipe at size 1, traded on the same dates at the same costs.",
    "",
    "Sheet Check: the column 'difference' should be zero to about 1e-15 (floating point).",
]


def _write(path: Path, cells: dict, values: dict | None = None) -> None:
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(path))
    fmt = {"date": wb.add_format({"num_format": "yyyy-mm-dd"}), "num": wb.add_format({"num_format": "0.000000000"}),
           "head": wb.add_format({"bold": True, "text_wrap": True, "valign": "top"})}
    sheets = {}
    readme = wb.add_worksheet("README")
    readme.set_column(0, 0, 130)
    for i, line in enumerate(NOTES):
        readme.write(i, 0, line, fmt["head"] if i == 0 else None)
    for name in ("Check", "Monthly", "Daily", "Inputs"):
        sheets[name] = wb.add_worksheet(name)
        sheets[name].set_row(0, 45, fmt["head"])
        sheets[name].freeze_panes(1, 1)
        sheets[name].set_column(0, 60, 14)
    sheets["Check"].set_column(0, 1, 16)
    sheets["Inputs"].set_column(0, 0, 18)
    for (sheet, ref), x in cells.items():
        ws = sheets[sheet]
        if isinstance(x, str) and x.startswith("="):
            v = None if values is None else values.get((sheet, ref))
            ws.write_formula(ref, x, fmt["num"], v if v is not None else 0)
        elif isinstance(x, datetime):
            ws.write_datetime(ref, x, fmt["date"])
        elif isinstance(x, (int, float, np.integer, np.floating)):
            ws.write_number(ref, float(x), fmt["num"])
        else:
            ws.write(ref, x)
    wb.close()


def _evaluate(path: Path) -> dict:
    """Compute every formula with the `formulas` package: {(sheet, "A1"): value}."""
    import formulas

    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for key, val in sol.items():
        if "!" not in key or ":" in key:
            continue
        sheet, ref = key.split("!")
        sheet = sheet.split("]")[-1].strip("'").title()
        x = np.asarray(val.value).ravel()[0]
        out[(sheet, ref.replace("$", ""))] = x if isinstance(x, str) else float(x)
    return out


def build(path: Path = PATH) -> pd.DataFrame:
    """Write the workbook, compute its formulas, and store the values. Returns the Check sheet."""
    daily = inputs()
    cells = _cells(daily)
    res = engine_run(daily)
    for r in range(2, 2 + len(CHECK_MONTHS) * len(QUANTITIES)):          # the engine's numbers, for display
        m, q = cells[("Check", f"A{r}")], cells[("Check", f"B{r}")]
        cells[("Check", f"D{r}")] = float(res.loc[pd.Period(m, "M"), q])
    _write(path, cells)
    values = _evaluate(path)
    _write(path, cells, values)
    return read(path)[1]


def read(path: Path = PATH) -> tuple[pd.DataFrame, pd.DataFrame]:
    """The daily inputs and the Check sheet, as stored in the workbook (formula values included)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    rows = list(wb["Daily"].iter_rows(min_row=2, max_col=5, values_only=True))
    daily = pd.DataFrame([r[2:5] for r in rows], columns=["A", "B", "rf"],
                         index=pd.DatetimeIndex([r[0] for r in rows], name="date"))
    check = pd.DataFrame(list(wb["Check"].iter_rows(min_row=2, max_col=5, values_only=True)),
                         columns=["month", "quantity", "spreadsheet", "engine", "difference"])
    return daily, check


if __name__ == "__main__":
    check = build()
    print(check.to_string())
    print(f"\nLargest |difference|: {check['difference'].abs().max():.2e}  ->  {PATH}")
