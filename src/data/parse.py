"""
Readers for each raw file format. Each returns data on the source's own
dates; putting everything on one calendar happens in build.py.
Returns come back as decimals (0.01 = 1%).
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd


def read_french(path: Path) -> pd.DataFrame:
    """Ken French daily CSV (zipped): free-text preamble, a header row starting with
    ',', YYYYMMDD rows in percent, then a copyright line. -99.99 / -999 mean missing."""
    with zipfile.ZipFile(path) as z:
        name = next(n for n in z.namelist() if n.lower().endswith(".csv"))
        lines = z.read(name).decode("latin-1").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith(","))
    body = []
    for line in lines[start + 1:]:
        if not line[:8].isdigit():
            break                       # end of the daily block
        body.append(line)
    df = pd.read_csv(io.StringIO("\n".join([lines[start], *body])), index_col=0)
    df.index = pd.to_datetime(df.index.astype(str), format="%Y%m%d")
    df.index.name = "date"
    df.columns = df.columns.str.strip()
    return df.mask(df <= -99.99) / 100.0


def read_aqr(path: Path, sheet: str, column: str = "USA") -> pd.Series:
    """One column of an AQR data-set sheet: preamble rows, a 'DATE' header row,
    then MM/DD/YYYY rows of decimal returns."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        rows = wb[sheet].iter_rows(values_only=True)
        header = next(r for r in rows if r and r[0] == "DATE")
        j = header.index(column)
        dates, values = [], []
        for r in rows:
            if r[0] is None or r[j] is None:
                continue
            dates.append(r[0] if isinstance(r[0], datetime) else datetime.strptime(str(r[0]), "%m/%d/%Y"))
            values.append(float(r[j]))
    finally:
        wb.close()
    return pd.Series(values, index=pd.DatetimeIndex(dates, name="date"), name=column).sort_index()


def read_yahoo(path: Path) -> pd.DataFrame:
    """CSV written by download.yahoo_csv: split-adjusted OHLC, adj_close (splits and
    dividends), volume, cash dividends and split ratios."""
    return pd.read_csv(path, index_col="date", parse_dates=["date"]).sort_index()


def read_cboe_vix(path: Path) -> pd.Series:
    df = pd.read_csv(path)
    df.columns = df.columns.str.strip().str.upper()
    s = pd.Series(df["CLOSE"].to_numpy(float),
                  index=pd.DatetimeIndex(pd.to_datetime(df["DATE"], format="%m/%d/%Y"), name="date"),
                  name="vix")
    return s.sort_index()


def read_damodaran_sp500(path: Path) -> pd.Series:
    """Annual S&P 500 return including dividends, by calendar year ('Returns by year' sheet)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb["Returns by year"].iter_rows(values_only=True)
        header = next(r for r in rows if r and r[0] == "Year")
        j = next(i for i, c in enumerate(header) if isinstance(c, str) and c.startswith("S&P 500"))
        out = {}
        for r in rows:
            if not isinstance(r[0], (int, float)) or r[j] is None:
                break                   # summary rows follow the yearly block
            out[int(r[0])] = float(r[j])
    finally:
        wb.close()
    return pd.Series(out, name="sp500_damodaran").rename_axis("year")
